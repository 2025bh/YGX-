from __future__ import annotations
import hashlib, io, json, os, re, secrets, shutil, sqlite3, threading, time, uuid, zipfile
from contextlib import contextmanager
from pathlib import Path
from urllib.parse import urlparse
from fastapi import FastAPI, Request, HTTPException, UploadFile, File, Form, Depends
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from .domain import normalize, calculate, differences, key, natural, parse_rows

ROOT=Path(__file__).resolve().parents[1]
REPO=ROOT.parents[1]
DATA=Path(os.environ.get('QUOTE_DATA_DIR',ROOT/'data')).resolve()
DATA.mkdir(parents=True,exist_ok=True)
(DATA/'assets').mkdir(exist_ok=True)
(DATA/'rules').mkdir(exist_ok=True)
DB=DATA/'review.sqlite3'
app=FastAPI(title='忆古轩报价审核台',docs_url=None,redoc_url=None)

def dumps(x): return json.dumps(x,ensure_ascii=False,separators=(',',':'))
def uid(): return uuid.uuid4().hex
def now(): return time.time()

@contextmanager
def db():
    con=sqlite3.connect(DB,timeout=20); con.row_factory=sqlite3.Row
    con.execute('PRAGMA foreign_keys=ON'); con.execute('PRAGMA busy_timeout=20000')
    try:
        yield con; con.commit()
    except Exception:
        con.rollback(); raise
    finally: con.close()

def initialize():
    with db() as c:
        c.executescript('''
        PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY,name TEXT UNIQUE,role TEXT,salt TEXT,password TEXT);
        CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY,user_id TEXT,expires REAL);
        CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY,name TEXT,customer_type TEXT,demo INTEGER DEFAULT 0,source_sha TEXT,source_asset TEXT,has_reference INTEGER DEFAULT 0,created REAL);
        CREATE TABLE IF NOT EXISTS rows(id TEXT PRIMARY KEY,project_id TEXT,ai TEXT,reference TEXT,adopted TEXT,status TEXT DEFAULT 'pending',version INTEGER DEFAULT 1,updated REAL,FOREIGN KEY(project_id) REFERENCES projects(id));
        CREATE TABLE IF NOT EXISTS snapshots(id TEXT PRIMARY KEY,project_id TEXT,kind TEXT,filename TEXT,sha TEXT,asset TEXT,content TEXT,created REAL);
        CREATE TABLE IF NOT EXISTS assets(id TEXT PRIMARY KEY,filename TEXT,path TEXT,media_type TEXT,project_id TEXT,created REAL);
        CREATE TABLE IF NOT EXISTS corrections(id TEXT PRIMARY KEY,row_id TEXT,project_id TEXT,user_id TEXT,request_id TEXT UNIQUE,before_value TEXT,after_value TEXT,error_type TEXT,reason TEXT,reading TEXT,pricing TEXT,evidence TEXT,created REAL,request_digest TEXT);
        CREATE TABLE IF NOT EXISTS rules(id TEXT PRIMARY KEY,correction_id TEXT,project_id TEXT,scope TEXT,status TEXT,content TEXT,created REAL);
        CREATE TABLE IF NOT EXISTS devices(id TEXT PRIMARY KEY,user_id TEXT,token_sha TEXT,label TEXT,last_seen REAL);
        CREATE TABLE IF NOT EXISTS jobs(id TEXT PRIMARY KEY,device_id TEXT,row_id TEXT,payload TEXT,status TEXT,result TEXT,created REAL);
        CREATE TABLE IF NOT EXISTS mutation_requests(id TEXT PRIMARY KEY,user_id TEXT,kind TEXT,digest TEXT,result TEXT);
        CREATE TABLE IF NOT EXISTS previews(project_id TEXT,frame_index INTEGER,pdf_asset TEXT,PRIMARY KEY(project_id,frame_index));
        CREATE TABLE IF NOT EXISTS pair_checks(row_id TEXT,source TEXT,PRIMARY KEY(row_id,source));
        CREATE TABLE IF NOT EXISTS pair_bindings(row_id TEXT PRIMARY KEY,reference_key TEXT);
        CREATE TABLE IF NOT EXISTS frame_checks(project_id TEXT,frame_key TEXT,signature TEXT,status TEXT,note TEXT,evidence TEXT,created REAL,PRIMARY KEY(project_id,frame_key));
        CREATE TABLE IF NOT EXISTS review_events(id TEXT PRIMARY KEY,row_id TEXT,project_id TEXT,kind TEXT,before_value TEXT,after_value TEXT,expected_after_version INTEGER,user_id TEXT,request_id TEXT UNIQUE,created REAL);
        CREATE TABLE IF NOT EXISTS drawing_inventory(project_id TEXT PRIMARY KEY,source_sha TEXT,complete INTEGER,frames TEXT);
        CREATE TABLE IF NOT EXISTS preview_images(project_id TEXT,frame_index INTEGER,image_asset TEXT,PRIMARY KEY(project_id,frame_index));
        ''')
        c.execute("INSERT OR IGNORE INTO users VALUES ('local-admin','本机管理员','admin','','')")
        if 'request_digest' not in {r[1] for r in c.execute('PRAGMA table_info(corrections)')}:
            c.execute('ALTER TABLE corrections ADD COLUMN request_digest TEXT')
        if 'row_id' not in {r[1] for r in c.execute('PRAGMA table_info(assets)')}:
            c.execute('ALTER TABLE assets ADD COLUMN row_id TEXT');c.execute('ALTER TABLE assets ADD COLUMN row_version INTEGER')

initialize()

@app.middleware('http')
async def guard(request: Request, call_next):
    if request.method not in {'GET','HEAD','OPTIONS'}:
        origin=request.headers.get('origin')
        if origin and urlparse(origin).netloc!=request.headers.get('host'):
            return JSONResponse({'detail':'来源地址不匹配，请从审核台页面操作'},403)
        if request.cookies.get('quote_session') and not origin and not request.headers.get('authorization'):
            return JSONResponse({'detail':'缺少页面来源校验'},403)
    response=await call_next(request)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['X-Frame-Options']='DENY'
    response.headers['Referrer-Policy']='same-origin'
    if request.url.path.startswith('/api'): response.headers['Cache-Control']='no-store'
    return response

def user(request: Request):
    token=request.cookies.get('quote_session','')
    with db() as c:
        r=c.execute('SELECT u.* FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token=? AND s.expires>?',(hashlib.sha256(token.encode()).hexdigest(),now())).fetchone()
    if not r: raise HTTPException(401,'请先登录审核台')
    return dict(r)

def master(u=Depends(user)):
    if u['role'] not in ('admin','master'): raise HTTPException(403,'需要老师傅审核权限')
    return u

def admin(u=Depends(user)):
    if u['role']!='admin': raise HTTPException(403,'需要管理员权限')
    return u

def session_response(u,request):
    token=secrets.token_urlsafe(40)
    with db() as c: c.execute('INSERT INTO sessions VALUES (?,?,?)',(hashlib.sha256(token.encode()).hexdigest(),u['id'],now()+3600*12))
    response=JSONResponse({'id':u['id'],'name':u['name'],'role':u['role']})
    response.set_cookie('quote_session',token,httponly=True,samesite='strict',secure=request.url.scheme=='https',max_age=43200)
    return response

@app.get('/api/health')
def health(): return {'ok':True,'version':'0.2.0'}

@app.post('/api/auth/local')
def local_login(request: Request):
    if request.client.host not in ('127.0.0.1','::1','testclient'): raise HTTPException(403,'本机快捷入口仅限主机；其他电脑请登录')
    if urlparse(request.url._url).hostname not in ('127.0.0.1','localhost','::1','testserver'): raise HTTPException(403,'请从localhost打开本机入口')
    return session_response({'id':'local-admin','name':'本机管理员','role':'admin'},request)

class Login(BaseModel):
    name:str=Field(min_length=1,max_length=80)
    password:str=Field(min_length=1,max_length=128)

attempts={}
@app.post('/api/auth/login')
def login(body: Login,request: Request):
    ip=request.client.host; stamps=[x for x in attempts.get(ip,[]) if x>now()-300]
    if len(stamps)>=8: raise HTTPException(429,'尝试次数过多，请5分钟后重试')
    with db() as c: u=c.execute('SELECT * FROM users WHERE name=?',(body.name,)).fetchone()
    check=hashlib.pbkdf2_hmac('sha256',body.password.encode(),bytes.fromhex(u['salt']) if u and u['salt'] else b'empty',200000).hex()
    if not u or not u['password'] or not secrets.compare_digest(check,u['password']):
        attempts[ip]=stamps+[now()]; raise HTTPException(401,'用户名或密码不正确')
    attempts.pop(ip,None)
    return session_response(dict(u),request)

@app.get('/api/auth/me')
def me(u=Depends(user)): return {k:u[k] for k in ('id','name','role')}

@app.post('/api/auth/logout')
def logout(request: Request):
    with db() as c: c.execute('DELETE FROM sessions WHERE token=?',(hashlib.sha256(request.cookies.get('quote_session','').encode()).hexdigest(),))
    r=JSONResponse({'ok':True}); r.delete_cookie('quote_session'); return r

class Account(Login): role:str='master'
@app.post('/api/users')
def account(body: Account,u=Depends(admin)):
    if body.role not in ('admin','master') or len(body.password)<8: raise HTTPException(422,'密码至少8位；角色只能是管理员或老师傅')
    salt=secrets.token_bytes(16)
    try:
        with db() as c: c.execute('INSERT INTO users VALUES (?,?,?,?,?)',(uid(),body.name,body.role,salt.hex(),hashlib.pbkdf2_hmac('sha256',body.password.encode(),salt,200000).hex()))
    except sqlite3.IntegrityError: raise HTTPException(409,'用户名已存在')
    return {'ok':True}

def project(c,pid):
    p=c.execute('SELECT * FROM projects WHERE id=?',(pid,)).fetchone()
    if not p: raise HTTPException(404,'项目不存在')
    result=dict(p);result['previews']={r['frame_index']:r['pdf_asset'] for r in c.execute('SELECT frame_index,pdf_asset FROM previews WHERE project_id=?',(pid,))}
    result['preview_images']={r['frame_index']:r['image_asset'] for r in c.execute('SELECT frame_index,image_asset FROM preview_images WHERE project_id=?',(pid,))}
    return result

def review_source(c,pid):
    snapshots='|'.join(r['sha'] for r in c.execute('SELECT sha FROM snapshots WHERE project_id=? ORDER BY created,id',(pid,)))
    p=project(c,pid);inventory=c.execute('SELECT source_sha,complete,frames FROM drawing_inventory WHERE project_id=?',(pid,)).fetchone()
    return hashlib.sha256(dumps([snapshots,p['source_sha'],{'source_sha':inventory['source_sha'],'frames':json.loads(inventory['frames'])} if inventory else None]).encode()).hexdigest()

def row_view(r,p):
    out=dict(r)
    for f in ('ai','reference','adopted'): out[f]=json.loads(out[f]) if out[f] else None
    for f in ('ai','reference','adopted'):
        if out[f] and out[f].get('frame_index') in p.get('previews',{}):out[f]={**out[f],'pdf':p['previews'][out[f]['frame_index']]}
    out['diffs']=differences(out['ai'],out['reference']) if p['has_reference'] else []
    out['adopted']=calculate(out['adopted'],p['customer_type'])
    if not out['adopted'].get('bounds'):
        peers=[json.loads(x['adopted']) for x in r_connection_rows(p['id'])]
        peers=[v for v in peers if v.get('figure')==out['adopted']['figure'] and v.get('bounds')]
        sources={(v.get('frame_index'),v.get('frame_handle')) for v in peers}
        if len(sources)==1 and peers:
            v=peers[0]
            out['adopted']={**out['adopted'],**{k:v[k] for k in ('bounds','frame_index','frame_handle','images','pdf') if k in v}}
            if v.get('frame_index') in p.get('previews',{}):out['adopted']['pdf']=p['previews'][v['frame_index']]
    # Presentation links use the same-version independent green-number inventory.
    # Frozen quote inputs and their original provenance are not rewritten.
    with db() as c:inventory=c.execute('SELECT * FROM drawing_inventory WHERE project_id=?',(p['id'],)).fetchone()
    if inventory and inventory['source_sha']==p['source_sha']:
        frames=json.loads(inventory['frames'])
        for field in ('ai','reference','adopted'):
            v=out[field]
            if not v:continue
            candidates=[f for f in frames if f['figure']==v['figure']]
            old=next((f for f in frames if f['index']==v.get('frame_index')),None)
            target=candidates[0] if len(candidates)==1 else old if old and old['figure']==v['figure'] else None
            if target:
                if target['index']!=v.get('frame_index'):v.pop('images',None);v.pop('pdf',None)
                v.update(bounds=target['bounds'],frame_index=target['index'],frame_handle=target.get('handle'))
                if target['index'] in p.get('previews',{}):v['pdf']=p['previews'][target['index']]
                if target['index'] in p.get('preview_images',{}):v['images']=[p['preview_images'][target['index']]]
            else:
                for k in ('bounds','images','pdf','frame_index','frame_handle','handles'):v.pop(k,None)
                v['note']=v.get('note','')+'；图号定位存在冲突，待复核，不能按旧图框定位'
    return out

def r_connection_rows(pid):
    with db() as c:return c.execute("SELECT adopted FROM rows WHERE project_id=? AND status!='superseded'",(pid,)).fetchall()

@app.get('/api/projects')
def projects(u=Depends(user)):
    with db() as c:
        result=[]
        for pp in c.execute('SELECT * FROM projects ORDER BY demo,created DESC').fetchall():
            p=project(c,pp['id']); rows=[row_view(r,p) for r in c.execute("SELECT * FROM rows WHERE project_id=? AND status!='superseded'",(p['id'],))]
            p.update(total=len(rows),reviewed=sum(r['status']=='resolved' for r in rows),differences=sum(bool(r['diffs']) for r in rows),high_risk=sum(r['adopted'].get('risk')=='高' for r in rows),thumbnail=next((r['adopted'].get('images',[None])[0] for r in rows if r['adopted'].get('images')),None))
            result.append(p)
    return result

@app.get('/api/projects/{pid}')
def details(pid:str,u=Depends(user)):
    with db() as c:
        p=project(c,pid)
        rows=[row_view(r,p) for r in c.execute("SELECT * FROM rows WHERE project_id=? AND status!='superseded'",(pid,))]
        p['rows']=sorted(rows,key=lambda r:natural((r['ai'] or r['reference'] or r['adopted'])['figure']))
        p['snapshots']=[dict(s) for s in c.execute('SELECT id,kind,filename,sha,created FROM snapshots WHERE project_id=? ORDER BY created',(pid,))]
    return p

def store_asset(content,filename,media_type,pid=None,c=None,row_id=None,row_version=None):
    aid=uid(); suffix=Path(filename).suffix.lower()
    path=DATA/'assets'/(aid+suffix); path.write_bytes(content)
    values=(aid,Path(filename).name,str(path.relative_to(DATA)),media_type,pid,now(),row_id,row_version)
    sql='INSERT INTO assets (id,filename,path,media_type,project_id,created,row_id,row_version) VALUES (?,?,?,?,?,?,?,?)'
    if c is not None: c.execute(sql,values)
    else:
        with db() as con: con.execute(sql,values)
    return aid

@app.get('/api/assets/{aid}')
def asset(aid:str,u=Depends(user)):
    with db() as c: a=c.execute('SELECT * FROM assets WHERE id=?',(aid,)).fetchone()
    if not a: raise HTTPException(404,'图片或图纸尚未准备')
    p=(DATA/a['path']).resolve()
    if not p.is_relative_to(DATA/'assets'): raise HTTPException(400,'附件路径异常')
    return FileResponse(p,media_type=a['media_type'],filename=a['filename'],content_disposition_type='inline')

async def upload_bytes(file,limit=30_000_000):
    content=await file.read(limit+1)
    if len(content)>limit: raise HTTPException(413,'文件过大，请拆分后导入')
    return content

def validated_import(content,name,kind='ai',confirm_spec_units=False):
    try:
        if name.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(x.file_size for x in z.infolist())>200_000_000: raise ValueError('表格解压后过大')
        rows,warnings=parse_rows(content,name,allow_no_figure=kind=='reference',confirm_spec_units=confirm_spec_units)
        for row in rows:calculate(row)
        if len(rows)>5000: raise ValueError('单批最多5000项，请拆分导入')
        return rows,warnings
    except (ValueError,TypeError,KeyError,zipfile.BadZipFile) as e: raise HTTPException(422,str(e))

@app.post('/api/import/preview')
async def preview(file:UploadFile=File(...),kind:str=Form('ai'),confirm_spec_units:bool=Form(False),u=Depends(master)):
    if kind not in ('ai','reference'):raise HTTPException(422,'导入类型无效')
    content=await upload_bytes(file); rows,warnings=validated_import(content,file.filename or '',kind,confirm_spec_units)
    return {'rows':rows,'count':len(rows),'warnings':warnings,'sha':hashlib.sha256(content).hexdigest()}

def add_snapshot(c,pid,kind,name,content,rows):
    sha=hashlib.sha256(content).hexdigest()
    existing=c.execute('SELECT id FROM snapshots WHERE project_id=? AND kind=? AND sha=? AND content=?',(pid,kind,sha,dumps(rows))).fetchone()
    if existing:return existing['id'],False
    aid=store_asset(content,name,'application/octet-stream',pid,c)
    sid=uid(); c.execute('INSERT INTO snapshots VALUES (?,?,?,?,?,?,?,?)',(sid,pid,kind,name,sha,aid,dumps(rows),now()))
    return sid,True

@app.post('/api/import')
async def import_quote(file:UploadFile=File(...),name:str=Form(''),customer_type:str=Form('C'),kind:str=Form('ai'),project_id:str=Form(''),request_id:str=Form(''),confirm_spec_units:bool=Form(False),u=Depends(master)):
    if kind not in ('ai','reference') or customer_type not in ('C','B'):raise HTTPException(422,'导入类型无效')
    content=await upload_bytes(file); rows,warnings=validated_import(content,file.filename or '',kind,confirm_spec_units)
    digest=hashlib.sha256(dumps([hashlib.sha256(content).hexdigest(),name,customer_type,kind,project_id,confirm_spec_units]).encode()).hexdigest()
    request_id=request_id or 'import-'+u['id']+'-'+digest
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        previous=c.execute('SELECT * FROM mutation_requests WHERE id=?',(request_id,)).fetchone()
        if previous:
            if previous['user_id']!=u['id'] or previous['digest']!=digest:raise HTTPException(409,'导入标识冲突，请重新核对')
            return {**json.loads(previous['result']),'duplicate':True}
        if project_id:
            p=project(c,project_id)
            if kind!='reference': raise HTTPException(409,'AI原稿已冻结，请创建新项目版本导入')
        else:
            if kind!='ai' or not name.strip():raise HTTPException(422,'新项目需要名称及AI独立原稿')
            project_id=uid(); c.execute('INSERT INTO projects (id,name,customer_type,created) VALUES (?,?,?,?)',(project_id,name[:120],customer_type,now())); p=project(c,project_id)
        sid,created=add_snapshot(c,project_id,kind,file.filename,content,rows)
        if not created:return {'id':project_id,'snapshot':sid,'duplicate':True}
        if kind=='ai':
            for row in rows:c.execute('INSERT INTO rows (id,project_id,ai,adopted,updated) VALUES (?,?,?,?,?)',(uid(),project_id,dumps(row),dumps(row),now()))
        else:
            ai_rows=c.execute('SELECT * FROM rows WHERE project_id=? AND ai IS NOT NULL',(project_id,)).fetchall()
            candidates={}
            for row in rows:candidates.setdefault(key(row),[]).append(row)
            used=set(); ai_counts={}
            for r in ai_rows: ai_counts[key(json.loads(r['ai']))]=ai_counts.get(key(json.loads(r['ai'])),0)+1
            for r in ai_rows:
                ai=json.loads(r['ai']); matches=candidates.get(key(ai),[])
                binding=c.execute('SELECT reference_key FROM pair_bindings WHERE row_id=?',(r['id'],)).fetchone()
                if binding:matches=candidates.get(tuple(json.loads(binding['reference_key'])),[])
                ref=matches[0] if len(matches)==1 and (binding or ai_counts[key(ai)]==1) and id(matches[0]) not in used else None
                if ref is not None: used.add(id(ref))
                if ref is not None and ref['figure']=='待配图':ref={**ref,'figure':ai['figure']}
                c.execute("UPDATE rows SET reference=?,status='pending',version=version+1 WHERE id=?",(dumps(ref) if ref else None,r['id']))
            # Prior unmatched rows are retained as history, excluded from current review.
            c.execute("UPDATE rows SET status='superseded',version=version+1 WHERE project_id=? AND ai IS NULL AND reference IS NOT NULL",(project_id,))
            for row in rows:
                if id(row) not in used:c.execute('INSERT INTO rows (id,project_id,reference,adopted,updated) VALUES (?,?,?,?,?)',(uid(),project_id,dumps(row),dumps(row),now()))
            c.execute('UPDATE projects SET has_reference=1 WHERE id=?',(project_id,))
        result={'id':project_id,'snapshot':sid,'count':len(rows),'warnings':warnings}
        c.execute('INSERT INTO mutation_requests VALUES (?,?,?,?,?)',(request_id,u['id'],'import',digest,dumps(result)))
    return result

class Correction(BaseModel):
    expected_version:int
    request_id:str=Field(min_length=8,max_length=100)
    adopted:dict
    status:str='resolved'
    error_type:str='尺寸错误'
    reason:str=Field(default='',max_length=8000)
    reading:str=Field(default='',max_length=8000)
    pricing:str=Field(default='',max_length=8000)
    evidence:list[dict]=Field(default_factory=list,max_length=20)
    scope:str='project'
    source:str|None=None

@app.post('/api/rows/{rid}/corrections')
def correct(rid:str,b:Correction,u=Depends(master)):
    fingerprint=hashlib.sha256(json.dumps(b.model_dump(),sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    if b.status not in ('resolved','waiting') or b.scope not in ('project','C','B','shared'):raise HTTPException(422,'审核状态或规则范围无效')
    try:
        excluded=bool(b.adopted.get('excluded')); row=normalize(b.adopted); row['excluded']=excluded
    except (ValueError,TypeError) as e:raise HTTPException(422,str(e))
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        old=c.execute('SELECT * FROM rows WHERE id=?',(rid,)).fetchone()
        if not old:raise HTTPException(404,'报价项不存在')
        previous=c.execute('SELECT * FROM corrections WHERE request_id=?',(b.request_id,)).fetchone()
        if previous:
            if previous['row_id']!=rid or previous['user_id']!=u['id'] or previous['request_digest']!=fingerprint:raise HTTPException(409,'保存标识冲突，请重新核对后保存')
            p=project(c,old['project_id']);return {'saved':True,'duplicate':True,'row':row_view(old,p)}
        if old['version']!=b.expected_version:raise HTTPException(409,'这一项刚被其他人修改，请刷新后对照再保存')
        p=project(c,old['project_id'])
        if b.source is not None:
            current=review_source(c,p['id'])
            if current!=b.source:raise HTTPException(409,'对照来源已换版，请重新核对本题')
        try:row=calculate(row,p['customer_type'])
        except ValueError as e:raise HTTPException(422,str(e))
        if b.status=='resolved' and not excluded:
            if row['figure']=='待配图' or row['figure'].startswith('待核'):raise HTTPException(422,'请先为本项确认CAD绿色图号；不确定时暂时待查')
            if row['unit']=='待确认':raise HTTPException(422,'请先确认本项计价单位')
            if row['calc_mode']=='pending':raise HTTPException(422,'请先确认本项计算方式')
            base=row.get('area') if row['unit']=='㎡' else row.get('total_length') if row['unit']=='米' else row.get('quantity')
            if base is None:raise HTTPException(422,'工程量尚未完整，请补测或暂存待测量')
            if p['customer_type']=='B' and row['calc_mode'] not in ('confirmed_rect','confirmed_length','confirmed_count','manual'):raise HTTPException(422,'请先确认B端计价口径')
        previous_data=json.loads(old['adopted'])
        business_fields=('figure','region','name','unit','height','length','width','thickness','quantity','material','price','risk','calc_mode','minimum_area')
        changed=any(previous_data.get(k)!=row.get(k) for k in business_fields) or bool(previous_data.get('excluded'))!=excluded
        if changed and not (b.reason.strip() or b.reading.strip()):raise HTTPException(422,'修改数据时请填写错误理由或图纸读法')
        if row['calc_mode']=='manual' and b.status=='resolved' and not excluded and not b.pricing.strip():raise HTTPException(422,'人工展开/扣除工程量需要填写计算说明')
        for e in b.evidence:
            a=c.execute('SELECT project_id,row_id,row_version FROM assets WHERE id=?',(e.get('asset'),)).fetchone()
            if not a or a['project_id']!=old['project_id']:raise HTTPException(422,'纠错证据不属于当前项目')
            if a['row_id'] and a['row_id']!=rid:raise HTTPException(422,'截图绑定的是其他部件，请切回原部件')
            if a['row_version'] and a['row_version']!=b.expected_version:raise HTTPException(422,'截图获取后该项已变化，请重新核实证据')
        if b.scope!='project' and (b.status!='resolved' or not b.reading.strip() or not b.pricing.strip()):raise HTTPException(422,'通用规则需要完成核实并填写读图与计算方法')
        cid=uid()
        c.execute('INSERT INTO corrections VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(cid,rid,p['id'],u['id'],b.request_id,old['adopted'],dumps(row),b.error_type,b.reason,b.reading,b.pricing,dumps(b.evidence),now(),fingerprint))
        c.execute('UPDATE rows SET adopted=?,status=?,version=version+1,updated=? WHERE id=?',(dumps(row),b.status,now(),rid))
        c.execute('INSERT INTO review_events VALUES (?,?,?,?,?,?,?,?,?,?)',(uid(),rid,p['id'],'correction',dumps(dict(old)),dumps(row),old['version']+1,u['id'],b.request_id,now()))
        rule_id=uid(); status='PROJECT_CASE' if b.scope=='project' else 'PENDING_ADAPTER'
        rule={'id':rule_id,'correction_id':cid,'project_id':p['id'],'figure':row['figure'],'component':row['name'],'scope':b.scope,'status':status,'reading':b.reading,'pricing':b.pricing,'reason':b.reason,'evidence':b.evidence,'before':json.loads(old['adopted']),'after':row,'created':now()}
        c.execute('INSERT INTO rules VALUES (?,?,?,?,?,?,?)',(rule_id,cid,p['id'],b.scope,status,dumps(rule),now()))
        new=c.execute('SELECT * FROM rows WHERE id=?',(rid,)).fetchone()
    # Database is authoritative. Mirroring failure stays visible and can be retried.
    synced=True
    try:(DATA/'rules'/f'{rule_id}.json').write_text(json.dumps(rule,ensure_ascii=False,indent=2),encoding='utf-8')
    except OSError:synced=False
    return {'saved':True,'row':row_view(new,p),'rule_status':status,'local_rule_file':synced}

@app.get('/api/rows/{rid}/history')
def history(rid:str,u=Depends(user)):
    with db() as c:
        rows=c.execute('SELECT c.*,u.name AS author FROM corrections c JOIN users u ON u.id=c.user_id WHERE row_id=? ORDER BY created DESC',(rid,)).fetchall()
    return [{**dict(r),**{f:json.loads(r[f]) for f in ('before_value','after_value','evidence')}} for r in rows]

class Pair(BaseModel): target_id:str; expected_version:int
@app.post('/api/rows/{rid}/pair')
def pair(rid:str,b:Pair,u=Depends(master)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        a=c.execute('SELECT * FROM rows WHERE id=?',(rid,)).fetchone(); r=c.execute('SELECT * FROM rows WHERE id=?',(b.target_id,)).fetchone()
        if not a or not r or a['project_id']!=r['project_id'] or not a['ai'] or r['ai'] or not r['reference'] or a['reference'] or r['status']=='superseded':raise HTTPException(422,'请选择同项目中尚未配对的人工报价项')
        if a['version']!=b.expected_version:raise HTTPException(409,'数据已变化，请刷新')
        ref=json.loads(r['reference'])
        if ref['figure']=='待配图':ref={**ref,'figure':json.loads(a['ai'])['figure']}
        c.execute("UPDATE rows SET reference=?,version=version+1,status='pending' WHERE id=?",(dumps(ref),rid)); c.execute("UPDATE rows SET status='superseded',version=version+1 WHERE id=?",(b.target_id,))
        c.execute('INSERT OR REPLACE INTO pair_bindings VALUES (?,?)',(rid,dumps(key(json.loads(r['reference'])))))
    return {'ok':True}

class NewRow(BaseModel): row:dict; reason:str=Field(min_length=1); request_id:str=''
@app.post('/api/projects/{pid}/rows')
def add_row(pid:str,b:NewRow,u=Depends(master)):
    try:r=normalize(b.row)
    except (ValueError,TypeError) as e:raise HTTPException(422,str(e))
    r['note']=b.reason
    with db() as c:
        c.execute('BEGIN IMMEDIATE');p=project(c,pid)
        digest=hashlib.sha256(dumps([pid,r]).encode()).hexdigest();req=b.request_id or 'add-'+u['id']+'-'+digest
        previous=c.execute('SELECT * FROM mutation_requests WHERE id=?',(req,)).fetchone()
        if previous:
            if previous['digest']!=digest or previous['user_id']!=u['id']:raise HTTPException(409,'补漏标识冲突')
            return {**json.loads(previous['result']),'duplicate':True}
        rid=uid(); c.execute('INSERT INTO rows (id,project_id,adopted,updated) VALUES (?,?,?,?)',(rid,pid,dumps(r),now()))
        c.execute('INSERT INTO mutation_requests VALUES (?,?,?,?,?)',(req,u['id'],'add',digest,dumps({'id':rid})))
    return {'id':rid}

@app.post('/api/projects/{pid}/evidence')
async def evidence(pid:str,file:UploadFile=File(...),row_id:str=Form(''),row_version:int=Form(0),u=Depends(master)):
    content=await upload_bytes(file,20_000_000)
    if content.startswith(b'\x89PNG\r\n\x1a\n'):suffix='.png';mime='image/png'
    elif content.startswith(b'\xff\xd8\xff'):suffix='.jpg';mime='image/jpeg'
    elif content.startswith(b'%PDF'):
        from pypdf import PdfReader
        try:PdfReader(io.BytesIO(content))
        except Exception:raise HTTPException(422,'PDF无法解析')
        suffix='.pdf';mime='application/pdf'
    else:raise HTTPException(422,'证据支持PNG、JPG或PDF')
    with db() as c:
        project(c,pid)
        if row_id:
            r=c.execute('SELECT project_id,version FROM rows WHERE id=?',(row_id,)).fetchone()
            if not r or r['project_id']!=pid or r['version']!=row_version:raise HTTPException(409,'报价项已变化，请刷新后重新附证据')
    aid=store_asset(content,'纠错证据'+suffix,mime,pid,row_id=row_id or None,row_version=row_version or None)
    return {'asset':aid,'mime':mime,'name':file.filename,'row_id':row_id,'row_version':row_version}

@app.get('/api/rules')
def rules(u=Depends(user)):
    with db() as c:return [json.loads(r['content']) for r in c.execute('SELECT content FROM rules ORDER BY created DESC')]

@app.get('/api/knowledge')
def knowledge(u=Depends(user)):
    path=registry_path()
    registry=json.loads(path.read_text(encoding='utf-8'))
    return {'version':registry.get('schema_version'),'rules':[{k:r.get(k) for k in ('rule_id','rule_name','status','business_explanation','calculation_method','layer','maturity')} for r in registry['rules'] if r.get('status')!='SUPERSEDED']}

def registry_path():
    override=os.environ.get('QUOTE_KNOWLEDGE_REGISTRY')
    if override:return Path(override)
    private=REPO/'报价系统/cad_quote_knowledge/00_core/RULE_REGISTRY.json'
    return private if private.exists() else ROOT/'assets/public_rules.json'

@app.post('/api/rules/sync')
def sync_rules(u=Depends(master)):
    with db() as c:rows=c.execute('SELECT id,content FROM rules').fetchall()
    for r in rows:(DATA/'rules'/f"{r['id']}.json").write_text(json.dumps(json.loads(r['content']),ensure_ascii=False,indent=2),encoding='utf-8')
    return {'saved':len(rows),'status':'已保存本地案例和待适配规则；主经验库不被静默覆盖'}

@app.get('/api/projects/{pid}/export')
def export(pid:str,u=Depends(user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.drawing.image import Image
    with db() as c:
        p=project(c,pid); rows=[row_view(r,p) for r in c.execute("SELECT * FROM rows WHERE project_id=? AND status!='superseded'",(pid,))]
        assets={r['id']:dict(r) for r in c.execute('SELECT * FROM assets WHERE project_id=?',(pid,))}
    rows.sort(key=lambda r:natural(r['adopted']['figure']))
    wb=Workbook();ws=wb.active;ws.title='报价审核稿'
    headers=['序号','区域','部件名称','参考图片','图号','单位','高度','长度','宽度','厚度','数量','总长','面积','材料','单价','小计','总计','风险等级']
    ws.append([p['name']+' · '+('演示审核稿' if p['demo'] else '报价审核稿')]);ws.merge_cells('A1:R1');ws.append(headers);ws.freeze_panes='G3'
    for cell in ws[2]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='2563EB')
    for i,r in enumerate(rows,3):
        v=r['adopted']; vals=[i-2,v.get('region'),v['name'],'待补CAD原图',v['figure'],v['unit'],v.get('height'),v.get('length'),v.get('width'),v.get('thickness'),v.get('quantity'),v.get('total_length'),v.get('area'),v.get('material'),v.get('price'),v.get('amount'),None,v.get('risk')]
        if v.get('excluded'):vals[11]=vals[12]=vals[15]=None;vals[2]+='（不计价）'
        ws.append(vals)
        # External descriptive text is always literal; only our generated quantity formulas execute.
        for col in (2,3,4,5,6,14,18):ws.cell(i,col).data_type='s'
        if not v.get('excluded') and v.get('calc_mode','') in ('rect','confirmed_rect') and v.get('area') is not None and not v.get('minimum_area'):ws.cell(i,13,f'=G{i}*H{i}*K{i}/1000000')
        if not v.get('excluded') and v.get('calc_mode','') in ('length','confirmed_length') and v.get('total_length') is not None:ws.cell(i,12,f'=H{i}*K{i}/1000')
        if v.get('amount') is not None:
            col='M' if v['unit']=='㎡' else 'L' if v['unit']=='米' else 'K';ws.cell(i,16,f'={col}{i}*O{i}')
        for j in (12,13,15,16,17):ws.cell(i,j).number_format='0.00'
        for cell in ws[i]:cell.alignment=Alignment(vertical='center',wrap_text=True)
        if r['status']!='resolved':ws.cell(i,3).fill=PatternFill('solid',fgColor='FFF3CD')
        image_id=next((a for a in v.get('images',[]) if a in assets and assets[a]['media_type'].startswith('image/')),None)
        if image_id:
            try:
                im=Image(DATA/assets[image_id]['path']);ratio=min(160/im.width,100/im.height);im.width*=ratio;im.height*=ratio;ws.add_image(im,f'D{i}');ws.cell(i,4,'');ws.row_dimensions[i].height=80
            except ImportError:ws.cell(i,4,'参考图片见平台')
    ws.cell(len(rows)+3,3,'木作合计（审核稿）')
    nonexcluded=[i for i,r in enumerate(rows,3) if not r['adopted'].get('excluded')]
    if nonexcluded and all(ws.cell(i,16).value is not None for i in nonexcluded):ws.cell(len(rows)+3,17,'='+'+'.join(f'P{i}' for i in nonexcluded))
    for letter,width in {'A':6,'B':22,'C':30,'D':24,'E':9,'F':7,'G':11,'H':11,'I':11,'J':9,'K':8,'L':11,'M':11,'N':44,'O':12,'P':14,'Q':14,'R':11}.items():ws.column_dimensions[letter].width=width
    buf=io.BytesIO();wb.save(buf);buf.seek(0)
    return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':"attachment; filename*=UTF-8''"+__import__('urllib.parse',fromlist=['quote']).quote(p['name']+'_审核稿.xlsx')})

@app.get('/api/backup')
def backup(u=Depends(admin)):
    backup_dir=DATA/'backups';backup_dir.mkdir(exist_ok=True); backup_db=backup_dir/(uid()+'.sqlite3')
    with db() as c:
        dst=sqlite3.connect(backup_db);c.backup(dst);dst.close()
    entries={'review.sqlite3':backup_db.read_bytes()}
    for folder in ('assets','rules'):
        for f in (DATA/folder).glob('*'):
            if f.is_file():entries[f.relative_to(DATA).as_posix()]=f.read_bytes()
    entries['knowledge_snapshot/RULE_REGISTRY.json']=registry_path().read_bytes()
    manifest={name:hashlib.sha256(content).hexdigest() for name,content in entries.items()}
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
        for name,content in entries.items():z.writestr(name,content)
        z.writestr('manifest.json',dumps({'schema':1,'sha256':manifest}))
    buf.seek(0)
    return StreamingResponse(buf,media_type='application/zip',headers={'Content-Disposition':'attachment; filename="quote-backup.zip"'})

@app.post('/api/devices')
def create_device(request:Request,u=Depends(master)):
    token=secrets.token_urlsafe(40); did=uid()
    with db() as c:c.execute('INSERT INTO devices VALUES (?,?,?,?,?)',(did,u['id'],hashlib.sha256(token.encode()).hexdigest(),'我的CAD电脑',0))
    config={'server':str(request.base_url).rstrip('/'),'token':token,'device':did}
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
        z.write(ROOT/'helper/CAD连接助手.ps1','CAD连接助手.ps1');z.write(ROOT/'helper/启动CAD连接.cmd','启动CAD连接.cmd');z.writestr('连接配置.json',json.dumps(config,ensure_ascii=False))
    buf.seek(0)
    return StreamingResponse(buf,media_type='application/zip',headers={'Content-Disposition':'attachment; filename="CAD-helper.zip"'})

def device(request:Request):
    token=request.headers.get('authorization','').removeprefix('Bearer ')
    with db() as c:
        r=c.execute('SELECT * FROM devices WHERE token_sha=?',(hashlib.sha256(token.encode()).hexdigest(),)).fetchone()
        if not r:raise HTTPException(401,'CAD连接未配对')
        c.execute('UPDATE devices SET last_seen=? WHERE id=?',(now(),r['id']))
    return dict(r)

@app.get('/api/devices')
def devices(u=Depends(master)):
    with db() as c:return [{**dict(r),'online':r['last_seen']>now()-20} for r in c.execute('SELECT id,label,last_seen FROM devices WHERE user_id=?',(u['id'],))]

class CADAction(BaseModel): action:str; device_id:str
@app.post('/api/rows/{rid}/cad')
def cad(rid:str,b:CADAction,u=Depends(master)):
    if b.action not in ('open','capture'):raise HTTPException(422,'不支持的CAD动作')
    with db() as c:
        d=c.execute('SELECT * FROM devices WHERE id=? AND user_id=?',(b.device_id,u['id'])).fetchone()
        if not d or d['last_seen']<now()-20:raise HTTPException(409,'CAD助手未连接，请先启动助手')
        r=c.execute('SELECT * FROM rows WHERE id=?',(rid,)).fetchone()
        if not r:raise HTTPException(404,'报价项不存在')
        p=project(c,r['project_id']);v=row_view(r,p)['adopted']
        if not p['source_asset'] or not v.get('bounds'):raise HTTPException(422,'当前项缺少同版DWG或图框定位范围，请先补充定位')
        job=uid();payload={'action':b.action,'project':p['id'],'source_asset':p['source_asset'],'source_sha':p['source_sha'],'review_source':review_source(c,p['id']),'bounds':v['bounds'],'figure':v['figure'],'component':v['name'],'handles':v.get('handles',[]),'row_version':r['version']}
        c.execute('INSERT INTO jobs VALUES (?,?,?,?,?,?,?)',(job,b.device_id,rid,dumps(payload),'queued',None,now()))
    return {'job':job}

@app.get('/api/cad/poll')
def poll(d=Depends(device)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        # A crashed helper leaves an explicit failure instead of replaying a capture on another view.
        c.execute("UPDATE jobs SET status='failed',result=? WHERE device_id=? AND status='running' AND created<?",(dumps({'message':'CAD助手任务超时，请重新定位后重试'}),d['id'],now()-120))
        j=c.execute("SELECT * FROM jobs WHERE device_id=? AND status='queued' ORDER BY created LIMIT 1",(d['id'],)).fetchone()
        if not j:return {'job':None}
        c.execute("UPDATE jobs SET status='running' WHERE id=?",(j['id'],))
    return {'job':j['id'],'payload':json.loads(j['payload'])}

@app.get('/api/cad/source/{aid}')
def cad_source(aid:str,d=Depends(device)):
    with db() as c:
        p=c.execute('SELECT * FROM projects WHERE source_asset=?',(aid,)).fetchone()
        active=c.execute("SELECT 1 FROM jobs WHERE device_id=? AND status='running' AND payload LIKE ?",(d['id'],'%'+aid+'%')).fetchone()
        a=c.execute('SELECT * FROM assets WHERE id=?',(aid,)).fetchone()
    if not p or not a or not active:raise HTTPException(403,'没有对应的CAD定位任务')
    return FileResponse(DATA/a['path'],media_type='application/octet-stream')

@app.post('/api/cad/jobs/{jid}/result')
async def cad_result(jid:str,status:str=Form(...),message:str=Form(''),file:UploadFile|None=File(None),d=Depends(device)):
    if status not in ('done','failed'):raise HTTPException(422,'CAD结果状态无效')
    content=await upload_bytes(file,20_000_000) if file else None
    if content is not None and not content.startswith(b'\x89PNG\r\n\x1a\n'):raise HTTPException(422,'CAD回传应为PNG')
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        j=c.execute("SELECT * FROM jobs WHERE id=? AND device_id=? AND status='running'",(jid,d['id'])).fetchone()
        if not j:raise HTTPException(409,'任务已结束或不属于此电脑')
        payload=json.loads(j['payload'])
        row=c.execute('SELECT * FROM rows WHERE id=?',(j['row_id'],)).fetchone()
        p=project(c,payload['project'])
        if (j['row_id'] and (not row or row['version']!=payload['row_version'] or row['status']=='superseded')) or p['source_sha']!=payload['source_sha'] or payload.get('review_source') and payload['review_source']!=review_source(c,p['id']):
            c.execute("UPDATE jobs SET status='failed',result=? WHERE id=?",(dumps({'message':'CAD取图期间数据或源图已变化，请重新定位'}),jid))
            return {'ok':False,'stale':True}
        result={'message':message[:1000],'row_id':row['id'] if row else None,'row_version':row['version'] if row else None,'source_sha':p['source_sha']}
        if content is not None:result['asset']=store_asset(content,'CAD当前视图.png','image/png',payload['project'],c,row['id'] if row else None,row['version'] if row else None)
        c.execute("UPDATE jobs SET status=?,result=? WHERE id=? AND status='running'",(status,dumps(result),jid))
    return {'ok':True}

@app.get('/api/cad/jobs/{jid}')
def job_status(jid:str,u=Depends(master)):
    with db() as c:j=c.execute('SELECT j.* FROM jobs j JOIN devices d ON d.id=j.device_id WHERE j.id=? AND d.user_id=?',(jid,u['id'])).fetchone()
    if not j:raise HTTPException(404,'CAD任务不存在')
    return {'status':j['status'],'result':json.loads(j['result']) if j['result'] else None}

# Built assets and SPA are served from the same origin as the API.
from . import review
if (ROOT/'dist/assets').exists():app.mount('/assets',StaticFiles(directory=ROOT/'dist/assets'),name='web-assets')
@app.get('/{path:path}')
def spa(path:str):
    if path.startswith('api/'):raise HTTPException(404,'接口不存在')
    index=ROOT/'dist/index.html'
    if not index.exists():raise HTTPException(503,'前端尚未构建，请先运行npm run build')
    return FileResponse(index)
