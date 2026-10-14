import io,json,os,uuid
from pathlib import Path
os.environ.setdefault('QUOTE_DATA_DIR',str(Path(__file__).resolve().parents[1]/'evidence'/('questions-'+uuid.uuid4().hex)))
import pytest
from fastapi.testclient import TestClient
from backend.app import app,db,DATA,store_asset
from backend.domain import normalize,calculate

@pytest.fixture
def client():
    c=TestClient(app,headers={'Origin':'http://testserver'})
    assert c.post('/api/auth/local',json={}).status_code==200
    return c

def row(**kw):
    return {'figure':'1','region':'模拟房间','name':'模拟柜','height':2100,'length':1200,'width':500,'quantity':1,'unit':'㎡','material':'模拟材质',**kw}

def imported(c,rows,ref=None,typ='C'):
    response=c.post('/api/import',data={'name':'模拟-'+uuid.uuid4().hex,'customer_type':typ,'kind':'ai'},files={'file':('ai.json',json.dumps({'rows':rows}).encode(),'application/json')})
    assert response.status_code==200,response.text
    pid=response.json()['id']
    if ref is not None:
        response=c.post('/api/import',data={'kind':'reference','project_id':pid},files={'file':('ref.json',json.dumps({'rows':ref}).encode(),'application/json')})
        assert response.status_code==200,response.text
    return pid,c.get('/api/projects/'+pid+'/review').json()

def answer(c,card,action='reference',**kw):
    payload={'action':action,'source':card['source'],'request_id':uuid.uuid4().hex,'expected_version':card['row']['version'],**kw}
    return c.post('/api/rows/'+card['id']+'/answer',json=payload)

def test_one_card_multiple_changes_and_freeze(client):
    pid,q=imported(client,[row(height=1900,length=1000)],[row(quantity=2)])
    assert q['pending']==1 and len(q['cards'][0]['changes'])==3
    a=answer(client,q['cards'][0]);assert a.status_code==200,a.text
    p=client.get('/api/projects/'+pid).json()
    assert p['rows'][0]['ai']['height']==1900 and p['rows'][0]['reference']['height']==2100
    assert p['rows'][0]['adopted']['area']==pytest.approx(5.04)

def test_pair_before_missing_extra_and_idempotent(client):
    pid,q=imported(client,[row(name='柜A')],[row(name='衣柜A')])
    assert len(q['cards'])==1 and q['cards'][0]['kind']=='pair'
    card=q['cards'][0];target=card['candidates'][0]
    a=answer(client,card,'pair',target_id=target['id'],target_version=target['version'])
    assert a.status_code==200,a.text
    q=client.get('/api/projects/'+pid+'/review').json()
    assert all(c['kind'] not in ('missing','extra','pair') for c in q['cards'])

def test_distinct_then_missing_and_extra(client):
    pid,q=imported(client,[row(name='柜A')],[row(name='脚线A')])
    assert answer(client,q['cards'][0],'distinct').status_code==200
    q=client.get('/api/projects/'+pid+'/review').json()
    assert {c['kind'] for c in q['cards']}=={'missing','extra'}

def test_region_not_cross_matched(client):
    pid,q=imported(client,[row(region='甲房'),row(region='乙房')],[row(region='乙房',height=2200),row(region='甲房',height=2000)])
    assert len(q['cards'])==2
    assert all(c['row']['ai']['region']==c['row']['reference']['region'] for c in q['cards'])

def test_linear_teacher_recalculates_without_area(client):
    pid,q=imported(client,[row(name='书桌',length=2800)],[row(name='书桌',length=2800,unit='米')])
    a=answer(client,q['cards'][0]);assert a.status_code==200,a.text
    assert a.json()['row']['adopted']['area'] is None
    assert a.json()['row']['adopted']['total_length']==2.8
    assert a.json()['row']['adopted']['amount'] is None

def test_custom_third_answer_and_pending(client):
    pid,q=imported(client,[row(height=1900)],[row(height=2100)])
    card=q['cards'][0]
    assert answer(client,card,'custom',adopted={'height':2050}).status_code==422
    a=answer(client,card,'custom',adopted={'height':2050},note='模拟完整标注高度2050')
    assert a.status_code==200 and a.json()['row']['adopted']['height']==2050
    pid,q=imported(client,[row(height=None)],[row(height=None)])
    # No difference still needs final whole-frame check, not fake completed.
    assert not q['finished']
    with db() as c:r=dict(c.execute('SELECT * FROM rows WHERE project_id=?',(pid,)).fetchone())
    card={'id':r['id'],'row':r,'source':q['source']}
    assert answer(client,card,'waiting',note='不知道答案，需补测').status_code==200

def test_blank_qty_is_not_default_one(client):
    pid,q=imported(client,[row(quantity=None)])
    assert q['cards'][0]['row']['adopted']['quantity'] is None
    assert answer(client,q['cards'][0],'ai').status_code==422

def test_undo_idempotency_and_frame_invalidation(client):
    pid,q=imported(client,[row(height=1900)],[row(height=2100)])
    a=answer(client,q['cards'][0]);assert a.status_code==200
    q=client.get('/api/projects/'+pid+'/review').json();f=q['frames'][0]
    assert client.post('/api/projects/'+pid+'/frame-review',json={'key':f['key'],'signature':f['signature']}).status_code==200
    r=a.json()['row'];body={'expected_version':r['version'],'request_id':uuid.uuid4().hex}
    url='/api/rows/'+r['id']+'/undo'
    assert client.post(url,json=body).status_code==200
    assert client.post(url,json=body).json()['duplicate']
    q=client.get('/api/projects/'+pid+'/review').json()
    assert q['frames'][0]['status']=='pending' and q['pending']==1
    assert q['cards'][0]['row']['adopted']['height']==1900

def test_answer_retry_and_changed_request_conflict(client):
    pid,q=imported(client,[row(height=1900)],[row()])
    card=q['cards'][0];body={'action':'reference','source':card['source'],'request_id':uuid.uuid4().hex,'expected_version':card['row']['version']}
    url='/api/rows/'+card['id']+'/answer'
    assert client.post(url,json=body).status_code==200
    assert client.post(url,json=body).json()['duplicate']
    assert client.post(url,json={**body,'note':'changed'}).status_code==409

def test_reference_replacement_reopens_and_preserves_manual_pair(client):
    pid,q=imported(client,[row(name='柜A')],[row(name='柜别称')]);card=q['cards'][0];t=card['candidates'][0]
    assert answer(client,card,'pair',target_id=t['id'],target_version=t['version']).status_code==200
    old=client.get('/api/projects/'+pid+'/review').json()['cards'][0]
    assert answer(client,old,'reference').status_code==200
    resp=client.post('/api/import',data={'kind':'reference','project_id':pid},files={'file':('new.json',json.dumps({'rows':[row(name='柜别称',height=2300)]}).encode(),'application/json')})
    assert resp.status_code==200
    q=client.get('/api/projects/'+pid+'/review').json()
    assert len(q['cards'])==1 and q['cards'][0]['row']['reference']['height']==2300
    assert answer(client,old,'ai').status_code==409

def test_frame_cannot_close_waiting_question(client):
    pid,q=imported(client,[row(height=None)])
    assert answer(client,q['cards'][0],'waiting').status_code==200
    q=client.get('/api/projects/'+pid+'/review').json();f=q['frames'][0]
    assert client.post('/api/projects/'+pid+'/frame-review',json={'key':f['key'],'signature':f['signature']}).status_code==422

def test_unknown_mode_and_b_pending(client):
    pid,q=imported(client,[row()],typ='B')
    assert answer(client,q['cards'][0],'ai').status_code==422
    assert answer(client,q['cards'][0],'custom',adopted={'calc_mode':'confirmed_length'},note='模拟错误单位匹配').status_code==422

def test_literal_excel_formula_text(client):
    pid,q=imported(client,[row(material='=1+1',name='=HYPERLINK("example")')])
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(client.get('/api/projects/'+pid+'/export').content),data_only=False)
    assert wb.active['C3'].data_type=='s' and wb.active['N3'].data_type=='s'

def test_equal_but_incomplete_still_needs_question(client):
    pid,q=imported(client,[row(height=None)],[row(height=None)])
    assert q['pending']==1 and q['cards'][0]['kind']=='verify'
    assert answer(client,q['cards'][0],'reference').status_code==422

def test_equal_frame_check_persists_and_confirms_rows(client):
    pid,q=imported(client,[row()],[row()])
    assert q['pending']==0 and not q['finished']
    f=q['frames'][0]
    assert client.post('/api/projects/'+pid+'/frame-review',json={'key':f['key'],'signature':f['signature']}).status_code==200
    after=client.get('/api/projects/'+pid+'/review').json()
    assert after['listed_complete'] and not after['finished'] and after['frames'][0]['status']=='done'
    assert after['frames'][0]['rows'][0]['status']=='resolved'

def test_drawing_same_version_and_geometry_validation(client):
    pid,q=imported(client,[row(bounds=[0,0,100,100])])
    a=client.post('/api/projects/'+pid+'/drawing',files={'file':('source.dwg',b'AC1032mock','application/octet-stream')})
    assert a.status_code==200
    assert client.post('/api/projects/'+pid+'/drawing',files={'file':('source.dwg',b'AC1032mock2','application/octet-stream')}).status_code==409
    assert client.post('/api/import/preview',files={'file':('bad.json',json.dumps({'rows':[row(bounds=[1,2,0,0])]}).encode(),'application/json')}).status_code==422

def test_cad_stale_result_and_bound_evidence(client):
    pid,q=imported(client,[row(bounds=[0,0,100,100])]);card=q['cards'][0]
    assert client.post('/api/projects/'+pid+'/drawing',files={'file':('source.dwg',b'AC1032mock','application/octet-stream')}).status_code==200
    card=client.get('/api/projects/'+pid+'/review').json()['cards'][0]
    import zipfile
    helper=client.post('/api/devices')
    z=zipfile.ZipFile(io.BytesIO(helper.content));cfg=json.loads(z.read('连接配置.json'))
    assert z.read('CAD连接助手.ps1').startswith(b'\xef\xbb\xbf')
    headers={'Authorization':'Bearer '+cfg['token']}
    client.get('/api/cad/poll',headers=headers)
    job=client.post('/api/rows/'+card['id']+'/cad',json={'action':'capture','device_id':cfg['device']}).json()['job']
    assert client.get('/api/cad/poll',headers=headers).json()['job']==job
    assert answer(client,card,'ai').status_code==200
    result=client.post('/api/cad/jobs/'+job+'/result',headers=headers,data={'status':'done'},files={'file':('cad.png',b'\x89PNG\r\n\x1a\nmock','image/png')})
    assert result.json()['stale']
    assert client.get('/api/cad/jobs/'+job).json()['status']=='failed'

def test_backup_hash_tamper_rejected(client,tmp_path):
    import zipfile
    from restore_backup import restore
    pid,q=imported(client,[row()])
    response=client.get('/api/backup');assert response.status_code==200
    original=zipfile.ZipFile(io.BytesIO(response.content));bad=tmp_path/'bad.zip'
    with zipfile.ZipFile(bad,'w') as z:
        for n in original.namelist():z.writestr(n,original.read(n)+(b'bad' if n=='review.sqlite3' else b''))
    with pytest.raises(ValueError):restore(bad,tmp_path/'restored')
    assert not (tmp_path/'restored').exists()

def test_pending_mode_never_resolved(client):
    pid,q=imported(client,[row(calc_mode='pending',price=10)])
    assert answer(client,q['cards'][0],'ai').status_code==422

def test_manual_source_quantity_preserved(client):
    pid,q=imported(client,[row(height=None,area=2.33)])
    assert q['cards'][0]['row']['adopted']['area']==2.33
    assert answer(client,q['cards'][0],'ai').status_code==200

def test_zero_quote_cad_frame_is_not_skipped(client):
    pid,q=imported(client,[row(bounds=[0,0,100,100],frame_index=1)],[row(bounds=[0,0,100,100],frame_index=1)])
    binding=client.post('/api/projects/'+pid+'/drawing',files={'file':('mock.dwg',b'AC1032mock','application/octet-stream')}).json()
    frames={'source_sha':binding['sha'],'complete':True,'frames':[{'index':1,'figure':'1','bounds':[0,0,100,100]},{'index':2,'figure':'2','bounds':[200,0,300,100]}]}
    assert client.post('/api/projects/'+pid+'/frames',json=frames).status_code==200
    q=client.get('/api/projects/'+pid+'/review').json()
    assert len(q['frames'])==2 and q['coverage_complete'] and not q['finished']
    zero=next(f for f in q['frames'] if not f['rows'])
    assert zero['drawing']['figure']=='2'
    payload={'key':zero['key'],'signature':zero['signature']}
    assert client.post('/api/projects/'+pid+'/frame-review',json=payload).status_code==422
    assert client.post('/api/projects/'+pid+'/frame-review',json={**payload,'note':'确认本图没有独立报价部件'}).status_code==200
    q=client.get('/api/projects/'+pid+'/review').json()
    assert not q['finished'] # the other frame still needs its own check

def test_hidden_pair_candidate_cannot_bypass_frame_validation(client):
    pid,q=imported(client,[row(name='甲柜',frame_index=1)],[row(name='乙柜',height=None,frame_index=2)])
    hidden=next(f for f in q['frames'] if any(not r['ai'] for r in f['rows']))
    assert client.post('/api/projects/'+pid+'/frame-review',json={'key':hidden['key'],'signature':hidden['signature']}).status_code==422

def test_reference_candidate_does_not_keep_old_price_or_floor(client):
    pid,q=imported(client,[row(height=1000,length=1000,price=100,minimum_area=2)],[row(height=1000,length=1000)])
    a=answer(client,q['cards'][0])
    assert a.status_code==200 and a.json()['row']['adopted']['area']==1 and a.json()['row']['adopted']['amount'] is None

def test_inventory_remaps_bad_link_without_rewriting_inputs(client):
    pid,q=imported(client,[row(frame_index=3,bounds=[0,0,100,100])])
    sha=client.post('/api/projects/'+pid+'/drawing',files={'file':('a.dwg',b'AC1032mock','application/octet-stream')}).json()['sha']
    old=client.get('/api/projects/'+pid+'/review').json()['cards'][0]
    assert client.post('/api/projects/'+pid+'/frames',json={'source_sha':sha,'frames':[{'index':3,'figure':'总览','bounds':[0,0,100,100]},{'index':4,'figure':'1','bounds':[200,0,300,100]}]}).status_code==200
    q=client.get('/api/projects/'+pid+'/review').json()
    assert q['cards'][0]['row']['adopted']['frame_index']==4
    with db() as c:
        original=json.loads(c.execute('SELECT ai FROM rows WHERE project_id=?',(pid,)).fetchone()['ai'])
        assert original['frame_index']==3
    assert answer(client,old,'ai').status_code==409

def test_teacher_sheet_without_figure_or_billing_unit(client):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    ws.append(['序号','区域','名称','单位','高度','长度','宽度','面积','单价'])
    ws.append([1,'模拟房间','模拟柜','mm',2100,1200,500,2.52,None])
    stream=io.BytesIO();wb.save(stream);content=stream.getvalue()
    p=client.post('/api/import/preview',data={'kind':'reference'},files={'file':('teacher.xlsx',content,'application/octet-stream')})
    assert p.status_code==200,p.text
    assert p.json()['rows'][0]['figure']=='待配图' and p.json()['rows'][0]['unit']=='待确认'
    p=client.post('/api/import/preview',data={'kind':'reference','confirm_spec_units':'true'},files={'file':('teacher.xlsx',content,'application/octet-stream')})
    assert p.status_code==200 and p.json()['rows'][0]['unit']=='㎡'
    pid,q=imported(client,[row()])
    r=client.post('/api/import',data={'project_id':pid,'kind':'reference','confirm_spec_units':'true'},files={'file':('teacher.xlsx',content,'application/octet-stream')})
    assert r.status_code==200,r.text
    card=client.get('/api/projects/'+pid+'/review').json()['cards'][0];assert card['kind']=='pair'
    target=card['candidates'][0]
    assert answer(client,card,'pair',target_id=target['id'],target_version=target['version']).status_code==200
    p=client.get('/api/projects/'+pid).json();assert p['rows'][0]['reference']['figure']=='1'

def test_display_rounding_does_not_turn_rectangle_into_manual():
    v=normalize(row(height=2101,length=1201,area=2.52))
    assert v['calc_mode']=='rect'
    assert calculate(v)['area']==pytest.approx(2.523301)

def test_excel_billing_unit_wins_over_spec_unit():
    from openpyxl import Workbook
    from backend.domain import parse_rows
    for headers,values in [(['单位','计价单位'],['mm','米']),(['计价单位','规格单位'],['米','mm'])]:
        wb=Workbook();ws=wb.active
        ws.append(['图号','部件名称','长度','数量']+headers)
        ws.append(['1','模拟脚线',3200,1]+values)
        stream=io.BytesIO();wb.save(stream)
        v=parse_rows(stream.getvalue(),'mock.xlsx',confirm_spec_units=True)[0][0]
        assert v['unit']=='米' and calculate(v)['total_length']==3.2
    assert normalize({**row(),'计价单位':'米','单位':'mm'})['unit']=='米'

def test_unlocated_teacher_missing_item_requires_figure(client):
    pid,q=imported(client,[row()],[{**row(figure='待配图',region='其他房间',name='漏项柜')}])
    card=next(c for c in q['cards'] if c['kind']=='missing')
    assert answer(client,card).status_code==422
    assert answer(client,card,'waiting',note='需核对绿色图号').status_code==200
    q=client.get('/api/projects/'+pid+'/review').json()
    card=next(c for c in q['cards'] if c['kind']=='missing')
    assert answer(client,card,'custom',adopted={**card['row']['adopted'],'figure':'2'},note='图2核实为这件漏项柜').status_code==200

def test_ambiguous_inventory_clears_unverified_bounds(client):
    pid,q=imported(client,[row(bounds=[0,0,100,100])])
    sha=client.post('/api/projects/'+pid+'/drawing',files={'file':('a.dwg',b'AC1032mock','application/octet-stream')}).json()['sha']
    frames=[{'index':3,'figure':'1','bounds':[0,0,100,100]},{'index':4,'figure':'1','bounds':[200,0,300,100]}]
    assert client.post('/api/projects/'+pid+'/frames',json={'source_sha':sha,'frames':frames}).status_code==200
    v=client.get('/api/projects/'+pid+'/review').json()['cards'][0]['row']['adopted']
    assert not v.get('bounds') and not v.get('frame_index')
